#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
════════════════════════════════════════════════════════════════════════════
 墨摘 · 微信公众号采集脚本 (wechat-digest skill) —— P0 全文挖掘与获取
 ─────────────────────────────────────────────────────────────────────────
 以「微信公众平台后台」的 Cookie+Token 接口为核心采集公众号文章：
   1) searchbiz  按公众号名称搜到 fakeid
   2) appmsg     按 fakeid 翻页拉取文章列表（标题/链接/摘要/发布时间/封面）
   3) 抓取每篇正文并清洗为纯文本，同时记录正文内图片 URL（全文挖掘）
   4) 落盘 output/articles_YYYYMMDD.json（供分析）+ index_YYYYMMDD.xlsx
   5) 默认自动「入库」到 output/knowledge_base.json（持续累积知识库）

 这是公众号无公开 API 时最稳定的采集路径，但仍受微信频控约束——脚本内置
 随机延时、指数退避、频控/凭证失效识别与进度保存，尽量稳。

 ── 命令（对标 redbook skill 的 search / read / whoami）──
   python3 wechat_collector.py collect 晚点LatePost --since 2025-01-01 --count 10
   python3 wechat_collector.py read  https://mp.weixin.qq.com/s/xxxx
   python3 wechat_collector.py whoami
   python3 wechat_collector.py                # 无参数 → 读文件顶部默认配置（旧用法）

 ⚠️ 阅读数/在看/点赞需微信客户端签名参数，后台接口拿不到，故不含互动量。
 ⚠️ token/cookie 会过期，采集失败优先重新获取（见 SKILL.md）。
════════════════════════════════════════════════════════════════════════════
"""

import argparse
import json
import os
import random
import re
import sys
import time
from datetime import datetime

import requests

# ════════════════════════════════════════════════════════════════
#  默认配置区（仅「无参数运行」时生效；推荐用命令行子命令）
# ════════════════════════════════════════════════════════════════

ACCOUNTS = [
    "晚点LatePost",
]
DATE_FILTER_THRESHOLD = "2025-01-01"   # 仅采此日期(含)之后的文章
TARGET_COUNT = 10                      # 每个号目标篇数
FETCH_CONTENT = True                   # 是否抓正文
AUTO_INGEST = True                     # 采完是否自动入知识库

# 输出目录
OUTPUT_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "output")
# 断点续采的游标状态（--resume 时读写；已随 output/ 一同 .gitignore）
COLLECT_STATE_FILE = os.path.join(OUTPUT_DIR, "collect_state.json")

# ── 反爬节流参数（频控很敏感，宁慢勿快）──
REQUEST_DELAY_MIN = 3.0
REQUEST_DELAY_MAX = 6.0
CONTENT_DELAY_MIN = 1.5
CONTENT_DELAY_MAX = 3.5
# 微信 appmsg 频控(ret=200013) 的实际冷却是几十分钟级，60s 后重试几乎必然再次撞墙；
# 这里默认拉到 5 分钟并逐次翻倍，真正的兜底是「已采集的进度一定落盘」。
FREQ_COOLDOWN = 300
FREQ_COOLDOWN_FACTOR = 2.0
MAX_FREQ_RETRIES = 2
LIST_PAGE_SIZE = 5
MAX_LIST_PAGES = 40
CONTENT_RETRIES = 3
CONTENT_BLOCKED_COOLDOWN = 30          # 文章页撞验证页时的等待（比列表接口轻）

UA = ("Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 "
      "(KHTML, like Gecko) Chrome/124.0.0.0 Safari/537.36")

MP_BASE = "https://mp.weixin.qq.com"
SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
CREDENTIALS_FILE = os.path.join(SCRIPT_DIR, "credentials.json")


# ════════════════════════════════════════════════════════════════
#  自定义异常
# ════════════════════════════════════════════════════════════════

class AuthError(Exception):
    """token / cookie 失效或未登录。"""


class FreqControlError(Exception):
    """命中微信频控，多次重试仍失败。"""


class ApiError(Exception):
    """接口返回非 0 的其它错误。"""


# ════════════════════════════════════════════════════════════════
#  规范化 id（与 kb.py.canonical_id 一致；优先复用 kb 模块）
# ════════════════════════════════════════════════════════════════

try:
    from kb import canonical_id  # type: ignore
except Exception:  # 容错：kb.py 缺失时本地降级
    import hashlib

    def canonical_id(link="", title=""):
        link = (link or "").strip()
        if link:
            m = re.search(r"[?&]sn=([0-9a-fA-F]+)", link)
            if m:
                return m.group(1).lower()
            return "h" + hashlib.sha1(link.encode("utf-8")).hexdigest()[:16]
        if title:
            return "t" + hashlib.sha1(title.encode("utf-8")).hexdigest()[:16]
        return "x" + hashlib.sha1(str(datetime.now()).encode("utf-8")).hexdigest()[:16]


# ════════════════════════════════════════════════════════════════
#  凭证加载
# ════════════════════════════════════════════════════════════════

def load_credentials():
    """优先读 credentials.json，其次环境变量 WECHAT_TOKEN / WECHAT_COOKIE。"""
    token = cookie = ""
    if os.path.exists(CREDENTIALS_FILE):
        try:
            with open(CREDENTIALS_FILE, "r", encoding="utf-8") as f:
                data = json.load(f)
            token = str(data.get("token", "")).strip()
            cookie = str(data.get("cookie", "")).strip()
        except Exception as e:
            print(f"⚠️  读取 credentials.json 失败：{e}")
    token = token or os.environ.get("WECHAT_TOKEN", "").strip()
    cookie = cookie or os.environ.get("WECHAT_COOKIE", "").strip()
    if token in ("在此粘贴 token 数字", "在此粘贴token数字"):
        token = ""
    if cookie in ("在此粘贴完整 cookie 字符串", "在此粘贴完整cookie字符串"):
        cookie = ""
    return token, cookie


def print_credential_guide():
    print("""
╭───────────────────────────────────────────────────────────────╮
│  未检测到有效的 token / cookie，无法采集。请按下面三步获取：     │
╰───────────────────────────────────────────────────────────────╯
  1. 浏览器登录 https://mp.weixin.qq.com （个人订阅号即可，免费）
  2. 登录后看地址栏 URL，形如：
       https://mp.weixin.qq.com/cgi-bin/home?...&token=1234567890
     末尾 token= 后的那串数字，就是 token。
  3. 按 F12 → Network 面板 → 刷新页面 → 点任意一个 mp.weixin.qq.com 的请求
     → Request Headers 里找到 Cookie，复制整条（很长一串）。

  然后把这两个值填进 skill/credentials.json：
       cp credentials.example.json credentials.json
       # 编辑 credentials.json，填入 token 和 cookie

  （或设置环境变量 WECHAT_TOKEN / WECHAT_COOKIE）

  ⚠️ token/cookie 会过期（通常几小时~几天），失效后重新获取即可。
""")


def make_session(cookie):
    s = requests.Session()
    s.headers.update({
        "User-Agent": UA,
        "Accept": "application/json, text/plain, */*",
        "Accept-Language": "zh-CN,zh;q=0.9,en;q=0.8",
        "Cookie": cookie,
        "Referer": MP_BASE + "/cgi-bin/appmsg",
    })
    return s


# ════════════════════════════════════════════════════════════════
#  正文清洗（移植自 backend/main.ts 的 htmlToText / sliceBetween / decodeEntities）
# ════════════════════════════════════════════════════════════════

def decode_entities(s):
    s = s or ""
    s = (s.replace("&amp;", "&").replace("&lt;", "<").replace("&gt;", ">")
          .replace("&quot;", '"').replace("&#39;", "'").replace("&nbsp;", " "))
    return s.strip()


def html_to_text(html):
    """把一段 HTML 转为带换行的纯文本（对标 main.ts 的 htmlToText）。"""
    if not html:
        return ""
    html = re.sub(r"<script[\s\S]*?</script>", "", html, flags=re.I)
    html = re.sub(r"<style[\s\S]*?</style>", "", html, flags=re.I)
    html = re.sub(r"<br\s*/?>", "\n", html, flags=re.I)
    html = re.sub(r"</(p|div|section|li|h[1-6])>", "\n", html, flags=re.I)
    html = re.sub(r"<[^>]+>", "", html)
    html = (html.replace("&nbsp;", " ").replace("&amp;", "&")
                .replace("&lt;", "<").replace("&gt;", ">")
                .replace("&quot;", '"').replace("&#39;", "'"))
    html = re.sub(r"&#x([0-9a-fA-F]+);", lambda m: chr(int(m.group(1), 16)), html)
    html = re.sub(r"&#(\d+);", lambda m: chr(int(m.group(1))), html)
    html = re.sub(r"[ \t]+\n", "\n", html)
    html = re.sub(r"\n{3,}", "\n\n", html)
    return html.strip()


def slice_between(html, start_re, end_re):
    """截取 start 正则之后、end 正则之前的片段（对标 main.ts 的 sliceBetween）。"""
    start_m = re.search(start_re, html, flags=re.I)
    if not start_m:
        return ""
    rest = html[start_m.end():]
    end_m = re.search(end_re, rest, flags=re.I)
    end_idx = end_m.start() if end_m else len(rest)
    return rest[:end_idx]


def _extract_div_inner(html, start_idx):
    """从 start_idx 处的 <div ...> 起做 div 深度配对，返回其内层 HTML。

    比「截到第一个 </div>」可靠：公众号正文里 div/section 多层嵌套，
    按第一个 </div> 截断会把正文腰斩；按「</div> 后面跟 <script」找又会越界
    多吞页面尾部。深度配对是唯一稳的做法。
    """
    open_end = html.find(">", start_idx)
    if open_end == -1:
        return ""
    pos = open_end + 1
    depth = 1
    for m in re.finditer(r"<div\b|</div\s*>", html[pos:], flags=re.I):
        if m.group(0).lower().startswith("</"):
            depth -= 1
            if depth == 0:
                return html[pos:pos + m.start()]
        else:
            depth += 1
    return html[pos:]


_BODY_TAG_RE = re.compile(r"<div\b[^>]*>", re.I)
_BODY_HIT_RE = re.compile(r'id\s*=\s*["\']js_content["\']'
                          r'|class\s*=\s*["\'][^"\']*rich_media_content[^"\']*["\']', re.I)


def slice_content_html(html):
    """取出 mp 文章页正文 div 的内层 HTML（未清洗，供取文本+图片）。

    一个页面可能有多个 id="js_content" / rich_media_content 容器——最典型的是
    「分享一篇文章」卡片页，它把 id="js_content" 挂在一个**空壳**
    <div class="original_panel_content"> 上，真正的正文在卡片指向的原文里。
    所以这里枚举所有候选容器、取「清洗后正文最长」的那个，空壳自然被淘汰。
    """
    if not html:
        return ""
    best, best_len = "", -1
    for m in _BODY_TAG_RE.finditer(html):
        if not _BODY_HIT_RE.search(m.group(0)):
            continue
        inner = _extract_div_inner(html, m.start())
        n = len(html_to_text(inner))
        if n > best_len:
            best, best_len = inner, n
    return best


def share_source_url(html):
    """「分享一篇文章」卡片页里「阅读全文」指向的原文链接（没有则空串）。"""
    m = re.search(r'<[a-z]+[^>]*id=["\']js_share_source["\'][^>]*>', html or "", flags=re.I)
    if not m:
        return ""
    u = re.search(r'data-url=["\']([^"\']+)["\']', m.group(0), flags=re.I)
    return decode_entities(u.group(1)).replace("&amp;", "&") if u else ""


def permanent_link(html, fallback=""):
    """从页面里的 biz/mid/idx 还原**永久链接**。

    搜狗给的 /s?src=11&timestamp=…&signature=… 链接会过期，且每次搜到的都不一样，
    直接拿它算 id 会导致同一篇文章反复入库。biz+mid+idx 唯一确定一篇文章。
    """
    def var(name):
        for m in re.finditer(r'var\s+%s\s*=\s*([^;]+);' % name, html or ""):
            for v in re.findall(r'"([^"]*)"', m.group(1)):
                if v:
                    return v
        return ""

    biz, mid, idx = var("biz"), var("mid") or var("appmsgid"), var("idx")
    if biz and mid:
        return "https://mp.weixin.qq.com/s?__biz=%s&mid=%s&idx=%s" % (biz, mid, idx or "1")
    return fallback


def extract_images(content_html):
    """从正文 HTML 提取图片 URL（微信图片懒加载在 data-src，常规在 src）。"""
    if not content_html:
        return []
    urls = []
    for m in re.finditer(r'<img[^>]+>', content_html, flags=re.I):
        tag = m.group(0)
        src = (re.search(r'data-src="([^"]+)"', tag, flags=re.I)
               or re.search(r"data-src='([^']+)'", tag, flags=re.I)
               or re.search(r'\ssrc="([^"]+)"', tag, flags=re.I)
               or re.search(r"\ssrc='([^']+)'", tag, flags=re.I))
        if not src:
            continue
        u = decode_entities(src.group(1)).replace("&amp;", "&")
        if u.startswith("//"):
            u = "https:" + u
        if u.startswith("http") and u not in urls:
            urls.append(u)
    return urls


def pick(html, re_):
    m = re.search(re_, html, flags=re.I)
    return decode_entities(re.sub(r"<[^>]+>", "", m.group(1))).strip() if m else ""


# ── 页面性质判定 ──────────────────────────────────────────────
# 微信对「失效链接 / 已删除 / 验证页」统统返回 HTTP 200 + 一个模板页，
# 只看状态码分辨不出来，必须看页面特征，否则只能得到「静默空正文」。
#
# 最关键的一点（实测得出）：微信真正的反爬页是 301/302 重定向到
#   https://mp.weixin.qq.com/mp/wappoc_appmsgcaptcha?poc_token=...&target_url=...
# 这个验证页的 <title> 是空的、正文区没有任何中文提示词，
# 所以只靠 HTML 内容根本认不出来——必须看**最终 URL**。
_BLOCKED_URL_RE = re.compile(r"appmsgcaptcha|verifycode|antispider|poc_token|/mp/captcha", re.I)
_BLOCKED_RE = re.compile(r"请输入验证码|antispider|环境异常|访问过于频繁|去验证"
                         r"|操作过于频繁|访问过快")
_GONE_RE = re.compile(r"已被发布者删除|因违规无法查看|内容已被投诉|无法查看此内容"
                      r"|该账号已自主注销|已被屏蔽|此内容发送失败")
_INVALID_RE = re.compile(r"参数错误|系统错误|请求参数有误|链接已过期")

STATUS_LABEL = {
    "ok": "",
    "digest": "正文未取到，已用摘要兜底",
    "empty": "页面正常但无正文（纯图片/视频推送）",
    "share": "分享卡片页，未能跟到原文",
    "blocked": "⚠ 被微信验证页拦截（需降速/换 IP 后重采）",
    "gone": "已删除 / 违规下架",
    "invalid": "⚠ 链接失效或参数错误（该链接已不可用）",
    "neterror": "⚠ 网络请求失败",
    "nolink": "列表未给出文章链接",
}


def classify_page(html, body_text="", final_url=""):
    """判定文章页性质，返回状态串。

    ok       正常取到正文
    blocked  命中验证页/反爬（可重试，但要降速）
    gone     已删除 / 违规 / 账号注销（永久失败，别重试）
    invalid  参数错误 / 链接失效（永久失败，别重试）
    share    「分享一篇文章」卡片页，正文在 data-url 指向的原文里
    empty    页面结构正常但确实没有正文（纯图片/视频推送）

    final_url 传 requests 的 r.url（跟完重定向后的地址）——这是识别
    微信验证页唯一可靠的信号，不传就会把风控误判成「链接失效」。
    """
    if final_url and _BLOCKED_URL_RE.search(final_url):
        return "blocked"
    if not html:
        return "invalid"
    if body_text and len(body_text) >= 20:
        return "ok"
    if _BLOCKED_RE.search(html):
        return "blocked"
    if _GONE_RE.search(html):
        return "gone"
    if share_source_url(html):
        return "share"
    if _INVALID_RE.search(html):
        return "invalid"
    if 'id="js_content"' not in html and "rich_media_content" not in html:
        # 连正文容器都没有：基本都是错误页 / 验证页模板
        return "invalid"
    return "empty"


def is_blocked_page(html):
    """向后兼容的薄封装（旧调用点仍可用）。"""
    return classify_page(html, "") == "blocked"


# ════════════════════════════════════════════════════════════════
#  节流 & HTTP 工具
# ════════════════════════════════════════════════════════════════

def polite_sleep(lo, hi):
    time.sleep(random.uniform(lo, hi))


def _as_int(v, default=0):
    """MP 接口的计数字段有时是 int 有时是字符串，统一成 int 再比较。

    不做这层归一化的话，`begin >= app_msg_cnt` 会抛
    TypeError: '>=' not supported between instances of 'int' and 'str'，
    把一整轮已经采好的文章连带 crash 掉。
    """
    try:
        return int(str(v).strip())
    except (TypeError, ValueError):
        return default


def api_get(session, token, path, extra_params):
    """调用 mp 后台 JSON 接口，带频控/凭证识别与重试。"""
    params = {"token": token, "lang": "zh_CN", "f": "json", "ajax": "1"}
    params.update(extra_params)
    url = MP_BASE + path

    for attempt in range(MAX_FREQ_RETRIES + 1):
        polite_sleep(REQUEST_DELAY_MIN, REQUEST_DELAY_MAX)
        try:
            r = session.get(url, params=params, timeout=20)
        except requests.RequestException as e:
            if attempt < MAX_FREQ_RETRIES:
                time.sleep(2 ** attempt + random.random())
                continue
            raise ApiError(f"网络请求失败：{e}")

        try:
            data = r.json()
        except ValueError:
            raise AuthError("接口未返回 JSON（通常是 token/cookie 失效，被重定向到登录页）。")

        ret = (data.get("base_resp") or {}).get("ret", 0)
        err = (data.get("base_resp") or {}).get("err_msg", "")

        if ret == 0:
            return data
        if ret == 200013:
            if attempt < MAX_FREQ_RETRIES:
                wait = int(FREQ_COOLDOWN * (FREQ_COOLDOWN_FACTOR ** attempt))
                print(f"   ⏳ 命中频控(ret=200013)，冷却 {wait}s 后重试"
                      f"（第 {attempt + 1}/{MAX_FREQ_RETRIES} 次）…")
                time.sleep(wait)
                continue
            raise FreqControlError(
                "多次重试仍被频控。微信 appmsg 的频控冷却通常是几十分钟到数小时，"
                "不是几秒——请隔一段时间再跑，并减小 --count / 减少账号数。")
        if ret in (200003, 200002, -3) or "session" in str(err).lower() or "token" in str(err).lower():
            raise AuthError(f"凭证失效或无权限（ret={ret}, err={err}）。请重新登录获取 token+cookie。")
        raise ApiError(f"接口错误 ret={ret}, err={err}")

    raise ApiError("未知错误")


def fetch_article(url, _depth=0):
    """
    抓取并清洗单篇正文，返回 dict：{content, images, ogCover, status, link}。

    status 见 classify_page —— 关键点是**把「空」的原因说清楚**：
    风控页 / 失效链接 / 已删除 / 真的没正文，这四种以前都被压成一个空串，
    导致「空了但没人知道为什么」，而且还会对永久失败的链接白重试 3 次。
    """
    headers = {
        "User-Agent": UA,
        "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
        "Accept-Language": "zh-CN,zh;q=0.9,en;q=0.8",
        "Referer": MP_BASE + "/",
    }
    last = {"content": "", "images": [], "ogCover": "", "status": "neterror", "link": url}
    for attempt in range(CONTENT_RETRIES):
        polite_sleep(CONTENT_DELAY_MIN, CONTENT_DELAY_MAX)
        try:
            r = requests.get(url, headers=headers, timeout=20)
            html = r.text or ""
            final_url = r.url or url
        except requests.RequestException:
            if attempt < CONTENT_RETRIES - 1:
                time.sleep(1.5 * (2 ** attempt) + random.random())
            continue

        content_html = slice_content_html(html)
        body = html_to_text(content_html) if content_html else ""
        images = extract_images(content_html)
        og_cover = pick(html, r'<meta property="og:image" content="([^"]*)"')
        status = classify_page(html, body, final_url)

        if status == "ok":
            return {"content": body, "images": images, "ogCover": og_cover,
                    "status": "ok", "link": permanent_link(html, url)}

        # 分享卡片页：正文在「阅读全文」指向的原文里，跟进一次
        if status == "share" and _depth == 0:
            src = share_source_url(html)
            if src:
                deep = fetch_article(src, _depth=1)
                if deep.get("content"):
                    deep["ogCover"] = deep.get("ogCover") or og_cover
                    return deep

        # 结构正常却取不到正文时，退回 og:description（只是摘要级兜底，如实标注）
        if status in ("empty", "share"):
            og = pick(html, r'<meta property="og:description" content="([^"]*)"')
            if og and len(og) > len(body):
                body, status = og, "digest"

        last = {"content": body, "images": images, "ogCover": og_cover,
                "status": status, "link": permanent_link(html, url)}

        # 永久性失败不必重试；只有风控页值得等更久再试一次
        if status != "blocked":
            break
        if attempt < CONTENT_RETRIES - 1:
            wait = CONTENT_BLOCKED_COOLDOWN * (attempt + 1)
            print(f"      ⏳ 文章页命中验证/风控，等 {wait}s 再试…")
            time.sleep(wait)
    return last


# ════════════════════════════════════════════════════════════════
#  采集逻辑
# ════════════════════════════════════════════════════════════════

def search_biz(session, token, name):
    """按名称搜公众号，返回 (fakeid, nickname)。找不到抛 ApiError。"""
    data = api_get(session, token, "/cgi-bin/searchbiz",
                   {"action": "search_biz", "begin": "0", "count": "5", "query": name})
    lst = data.get("list") or []
    if not lst:
        raise ApiError(f"未搜到公众号「{name}」（请确认名称是否精确）。")
    exact = next((b for b in lst if b.get("nickname") == name), None)
    chosen = exact or lst[0]
    if len(lst) > 1:
        print(f"   ℹ️  「{name}」匹配到 {len(lst)} 个号，已选用："
              f"{chosen.get('nickname')}（如需其它请调整名称）：")
        for b in lst[:5]:
            mark = "→" if b is chosen else " "
            print(f"      {mark} {b.get('nickname')}  (fakeid={b.get('fakeid')})")
    return chosen.get("fakeid"), chosen.get("nickname", name)


def unix_to_date(ts):
    try:
        return datetime.fromtimestamp(int(ts)).strftime("%Y-%m-%d")
    except Exception:
        return ""


def fetch_contents(records, label=""):
    """为一批记录补抓正文，并汇报真实的「空率」与空的原因分布。

    已有 content 的记录会跳过，所以中断后补抓、重复调用都安全。
    """
    if not records:
        return 0
    print(f"   ↓ 抓取正文（{len(records)} 篇）{label}…")
    okc = 0
    reasons = {}
    for i, rec in enumerate(records, 1):
        # 已有正文的跳过但**仍计为成功**，否则重跑/补抓时空率会被虚报。
        if rec.get("content"):
            okc += 1
            continue
        if not rec.get("link"):
            reasons["nolink"] = reasons.get("nolink", 0) + 1
            continue
        res = fetch_article(rec["link"])
        rec["content"] = res["content"]
        rec["images"] = res["images"]
        rec["contentStatus"] = res.get("status", "")
        if not rec.get("cover") and res.get("ogCover"):
            rec["cover"] = res["ogCover"]
        if res.get("status") == "ok":
            okc += 1
            flag = f"{len(res['content'])} 字 / {len(res['images'])} 图"
        else:
            st = res.get("status", "")
            reasons[st] = reasons.get(st, 0) + 1
            flag = f"未取到 · {STATUS_LABEL.get(st, st)}"
        print(f"      {i}/{len(records)} {flag}")

    n = len(records)
    print(f"   ↳ 正文成功 {okc}/{n}（空率 {(n - okc) / n * 100:.0f}%）")
    for st, c in sorted(reasons.items(), key=lambda kv: -kv[1]):
        print(f"      · {STATUS_LABEL.get(st, st)} × {c}")
    if reasons.get("blocked"):
        print("      → 出现风控页：请过一段时间对同一批重跑，"
              "kb ingest 会自动用更长的正文补全（不会重复入库）。")
    return okc


def collect_account(session, token, name, since, count, fetch_content, sink=None, resume=False):
    """采集单个公众号，返回文章记录列表。

    sink 若传入，则每采到一条就立即 append 进去。这是为了避开一个很痛的坑：
    以前频控/凭证失效是直接从函数里抛异常，调用方的 `all.extend(collect_account(...))`
    根本执行不到，于是屏幕上刚刚已经列印出来的几十篇全部丢弃，
    最后只给一句「未采集到任何文章」——这就是最主要的「空」来源。

    resume=True 时做**断点续采**：从上次存的 begin 游标接着翻，不重复请求已拉过的页——
    因为每翻一页 = 一次珍贵的 appmsg 请求，重复拉旧页等于白烧配额。
    """
    print(f"\n📰 采集公众号：{name}")
    fakeid, nickname = search_biz(session, token, name)
    print(f"   ✓ fakeid={fakeid}  昵称={nickname}")

    collected = []
    begin = load_collect_state(fakeid) if resume else 0
    if resume and begin:
        print(f"   ↻ 断点续采：从 begin={begin} 继续（跳过已拉位置，不重复花 appmsg 配额）")
    pages = 0
    now_str = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    try:
        while len(collected) < count and pages < MAX_LIST_PAGES:
            data = api_get(session, token, "/cgi-bin/appmsg", {
                "action": "list_ex", "begin": str(begin), "count": str(LIST_PAGE_SIZE),
                "fakeid": fakeid, "type": "9", "query": "",
            })
            items = data.get("app_msg_list") or []
            total = _as_int(data.get("app_msg_cnt"))
            if not items:
                break

            stop = False
            for it in items:
                date = unix_to_date(it.get("update_time"))
                if date and since and date < since:
                    stop = True
                    continue
                link = (it.get("link", "") or "").replace("\\/", "/")
                title = decode_entities(it.get("title", ""))
                rec = {
                    "id": canonical_id(link, title),
                    "account": nickname,
                    "query": name,
                    "title": title,
                    "digest": decode_entities(it.get("digest", "")),
                    "link": link,
                    "cover": it.get("cover", ""),
                    "images": [],
                    "publishDate": date,
                    "content": "",
                    "collectedAt": now_str,
                }
                collected.append(rec)
                if sink is not None:
                    sink.append(rec)
                print(f"   [{len(collected)}] {date}  {title[:36]}")
                if len(collected) >= count:
                    break

            pages += 1
            begin += LIST_PAGE_SIZE
            if resume:
                save_collect_state(fakeid, begin, nickname, since)
            if stop or (total and begin >= total):
                break
    except (AuthError, FreqControlError):
        # 列表中断也要把已拿到的这批正文补齐再抛，避免白跑一场。
        # 文章页走的是公开 mp 页面，与后台接口的频控是两套系统，通常仍可用。
        if resume:
            # 保住游标：begin 此刻指向未拿到的那页，下次 --resume 从这里重试。
            save_collect_state(fakeid, begin, nickname, since)
        if fetch_content and collected:
            print(f"   ⚠️ 列表中断，先把已采到的 {len(collected)} 篇正文抓完…")
            fetch_contents(collected, label="（中断前补抓）")
        raise

    if fetch_content and collected:
        fetch_contents(collected)

    print(f"   ✓ {name} 共采集 {len(collected)} 篇"
          + (f"（游标 begin={begin}）" if resume else ""))
    return collected


# ════════════════════════════════════════════════════════════════
#  落盘 & 入库
# ════════════════════════════════════════════════════════════════

def load_collect_state(fakeid):
    """读回该公众号（按 fakeid）上次翻到的 begin 游标；没有则 0。"""
    if not os.path.exists(COLLECT_STATE_FILE):
        return 0
    try:
        with open(COLLECT_STATE_FILE, "r", encoding="utf-8") as f:
            data = json.load(f)
        return int((data.get(fakeid) or {}).get("begin", 0))
    except (ValueError, OSError, TypeError):
        return 0


def save_collect_state(fakeid, begin, nickname, since):
    """把该公众号的 begin 游标存回状态文件（按 fakeid 分条，互不影响）。"""
    os.makedirs(OUTPUT_DIR, exist_ok=True)
    data = {}
    if os.path.exists(COLLECT_STATE_FILE):
        try:
            with open(COLLECT_STATE_FILE, "r", encoding="utf-8") as f:
                data = json.load(f)
        except (ValueError, OSError):
            data = {}
    data[fakeid] = {
        "begin": int(begin), "account": nickname, "since": since,
        "updatedAt": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
    }
    with open(COLLECT_STATE_FILE, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)


def save_json(articles, accounts, since, partial=False):
    os.makedirs(OUTPUT_DIR, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d")
    suffix = "_partial" if partial else ""
    path = os.path.join(OUTPUT_DIR, f"articles_{stamp}{suffix}.json")
    payload = {
        "generatedAt": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "dateFilter": since,
        "accounts": accounts,
        "count": len(articles),
        "partial": partial,
        "articles": articles,
    }
    with open(path, "w", encoding="utf-8") as f:
        json.dump(payload, f, ensure_ascii=False, indent=2)
    return path


def save_excel(articles, partial=False):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        print("⚠️  未安装 openpyxl，跳过 Excel（pip install openpyxl）。")
        return None

    stamp = datetime.now().strftime("%Y%m%d")
    suffix = "_partial" if partial else ""
    path = os.path.join(OUTPUT_DIR, f"index_{stamp}{suffix}.xlsx")

    wb = Workbook()
    ws = wb.active
    ws.title = "公众号文章索引"
    headers = ["序号", "公众号", "标题", "发布日期", "摘要", "正文字数", "图片数", "链接", "采集时间"]
    ws.append(headers)

    head_fill = PatternFill("solid", fgColor="1F6FEB")
    head_font = Font(bold=True, color="FFFFFF", size=11)
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill = head_fill
        cell.font = head_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for i, a in enumerate(articles, 1):
        ws.append([
            i, a.get("account", ""), a.get("title", ""), a.get("publishDate", ""),
            a.get("digest", ""), len(a.get("content", "") or ""),
            len(a.get("images", []) or []), a.get("link", ""), a.get("collectedAt", ""),
        ])

    widths = [6, 18, 48, 12, 50, 9, 7, 60, 20]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"

    wb.save(path)
    return path


def ingest_to_kb(articles):
    """采完自动入知识库（失败不影响已落盘的 JSON/Excel）。"""
    try:
        import kb
        store = kb.load_kb()
        added, updated = kb.ingest_articles(store, articles)
        kb.save_kb(store)
        print(f"   ↳ 已入知识库：新增 {added} 篇，补全 {updated} 篇，"
              f"库内共 {len(store['articles'])} 篇。")
        return True
    except Exception as e:
        print(f"   ⚠️ 入库失败（已保留 JSON/Excel，可稍后 `python3 kb.py ingest`）：{e}")
        return False


# ════════════════════════════════════════════════════════════════
#  子命令：search（按关键词搜全网公众号文章）
# ════════════════════════════════════════════════════════════════

def run_search(keyword, count, fetch_content, auto_ingest):
    """按关键词搜索微信公众号文章（跳号），抓取正文并入库。

    ⚠ 已用真实有效凭证验证：本命令依赖的
    `/cgi-bin/operate_appmsg?action=search_article` 返回
    **ret=200009, err_msg="not found"**，即该 action 并不存在。
    微信公众平台后台未提供「全网文章搜索」能力，因此本命令**永远返回空**。
    保留它只为向后兼容与错误可见，请改用 `collect <公众号名>`。
    """
    print("=" * 64)
    print("  墨摘 · 微信公众号文章搜索（关键词 → 跳号采集）")
    print("=" * 64)
    print("⚠️  该子命令已经真实凭证验证为**不可用**：微信后台的"
          "search_article 接口返回 ret=200009 (not found)。")
    print("    微信公众平台未提供「全网文章搜索」能力，本命令不会有结果。")
    print("    → 请改用：python3 wechat_collector.py collect <公众号名>")

    token, cookie = load_credentials()
    if not token or not cookie:
        print_credential_guide()
        return 1

    session = make_session(cookie)
    print(f"\n🔍 搜索关键词：{keyword}  目标 {count} 篇")

    collected = []
    begin = 0
    now_str = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    while len(collected) < count and begin < count + 50:
        try:
            data = api_get(session, token, "/cgi-bin/operate_appmsg", {
                "action": "search_article", "begin": str(begin),
                "count": str(min(5, count - len(collected))),
                "query": keyword, "type": "9",
            })
        except AuthError as e:
            print(f"✗ {e}")
            break
        except FreqControlError as e:
            print(f"✗ {e}")
            break
        except ApiError as e:
            print(f"✗ {e}")
            if "200009" in str(e):
                print("  → ret=200009 (not found) 即「该 action 不存在」，不是网络或频控问题，"
                      "重试也不会有结果。请改用 collect <公众号名>。")
            break

        items = data.get("article_list") or data.get("app_msg_list") or []
        if not items:
            # try alternate response key
            items = (data.get("results") or {}).get("list") or []
        if not items:
            # 这三个键名都是猜的，接口的 action 与响应结构从未被验证过。
            # 与其静默打一句「无更多结果」把问题藏起来，不如把真实响应摊开，
            # 让使用者/维护者一眼看出到底是「真没结果」还是「字段名不对」。
            if begin == 0:
                print("   ⚠️  未能从响应中解析出文章列表。")
                print(f"      响应顶层键名：{sorted(data.keys())}")
                sample = {k: v for k, v in data.items() if k != "base_resp"}
                dump = json.dumps(sample, ensure_ascii=False)[:600]
                print(f"      响应样本（截断）：{dump}")
                print("      说明：search 走的 operate_appmsg?action=search_article 属未验证接口，"
                      "微信后台并未公开「全网文章搜索」能力，此路径可能本就不可用。")
                print("      可靠替代：先用 collect <公众号名> 按号采集。")
            else:
                print("   (无更多结果)")
            break

        for it in items:
            link = (it.get("link") or it.get("url") or "").replace("\\/", "/")
            title = decode_entities(it.get("title") or "")
            account = decode_entities(it.get("nickname") or it.get("account_name") or "")
            date = ""
            ts = it.get("update_time") or it.get("publish_time") or it.get("create_time")
            if ts:
                date = unix_to_date(ts)
            collected.append({
                "id": canonical_id(link, title),
                "account": account,
                "query": keyword,
                "title": title,
                "digest": decode_entities(it.get("digest") or ""),
                "link": link,
                "cover": it.get("cover") or it.get("image_url") or "",
                "images": [],
                "publishDate": date,
                "content": "",
                "collectedAt": now_str,
            })
            print(f"   [{len(collected)}] {account} · {date}  {title[:36]}")
            if len(collected) >= count:
                break

        begin += 5

    if not collected:
        print(f"\n✗ 未搜到关于「{keyword}」的文章。")
        return 1

    # fetch full content
    if fetch_content and collected:
        fetch_contents(collected)

    # save
    json_path = save_json(collected, [keyword], "")
    xlsx_path = save_excel(collected)

    print("\n" + "=" * 64)
    print(f"✅ 完成：关键词「{keyword}」共 {len(collected)} 篇")
    print(f"   JSON ：{json_path}")
    if xlsx_path:
        print(f"   Excel：{xlsx_path}")
    if auto_ingest:
        ingest_to_kb(collected)
    print("=" * 64)
    return 0


# ════════════════════════════════════════════════════════════════
#  子命令：collect / read / whoami
# ════════════════════════════════════════════════════════════════

def print_fallback_hint(accounts=None):
    """collect（C/appmsg）拿不到数据时的引导：重申 C 为首选，把 A 作为需确认的备选。

    不自动切到 A：appmsg 是最直接最准确的路径（按真实文章列表拉），
    而 collect-urls 的完整性只取决于所给来源，可能漏文章/串同名号，需用户拍板才用。
    """
    print("\n── 采集路径优先级 ─────────────────────────────────")
    print("  首选（默认）：collect（appmsg 后台接口）——最直接、最准确，按公众号真实文章列表拉全文。")
    print("  当前是被微信频控挡住（不是数据没了）。建议做法按序：")
    print("    1) 等冷却 / 换个时段重跑 collect；拉大号的历史时加 --resume（断点续采，不重花配额）；")
    print("    2) 实在持续拿不到，且你手头有该号的文章 URL / RSS 源时，再【经你确认】改用备选：")
    print("         python3 wechat_collector.py collect-urls <URL / 文件 / RSS 源>")
    print("  ⚠ 备选路径（collect-urls）的文章完整性只取决于你给的来源，可能漏文章、甚至串到同名号，")
    print("    不如 collect 权威——所以默认不会自动切过去，由你决定。")


def run_collect(accounts, since, count, fetch_content, auto_ingest, resume=False):
    print("=" * 64)
    print("  墨摘 · 微信公众号采集（Cookie+Token）")
    print("=" * 64)

    token, cookie = load_credentials()
    if not token or not cookie:
        print_credential_guide()
        return 1
    if not accounts:
        print("⚠️  未指定公众号。用法：collect <名称...> 或在脚本顶部配置 ACCOUNTS。")
        return 1

    print(f"配置：{len(accounts)} 个公众号 | 日期≥{since or '不限'} | "
          f"每号目标 {count} 篇 | 抓正文={fetch_content} | 自动入库={auto_ingest}"
          + (" | 断点续采=开" if resume else ""))

    session = make_session(cookie)
    all_articles = []
    interrupted = False

    for name in accounts:
        try:
            # sink=all_articles：采到一篇就落入 all_articles，不再靠返回值。
            # 一旦抛频控/凭证异常，.extend() 是执行不到的，那就会把
            # 屏幕上已经列印出来的文章全部丢掉——这是之前「空」的根因。
            collect_account(session, token, name, since, count, fetch_content,
                            sink=all_articles, resume=resume)
        except AuthError as e:
            print(f"\n❌ 凭证问题：{e}")
            print_credential_guide()
            interrupted = True
            break
        except FreqControlError as e:
            print(f"\n⏸  {e}")
            interrupted = True
            break
        except ApiError as e:
            print(f"   ⚠️ 跳过「{name}」：{e}")
            continue

    if not all_articles:
        print("\n未采集到任何文章。请检查公众号名称、凭证有效性或日期阈值。")
        if interrupted:
            # 一篇都没拿到且是被频控中断：把首选 C / 备选 A（需确认）的处置说清楚。
            print_fallback_hint(accounts)
        return 1

    json_path = save_json(all_articles, accounts, since, partial=interrupted)
    xlsx_path = save_excel(all_articles, partial=interrupted)

    print("\n" + "=" * 64)
    print(f"✅ 完成{'（部分，已中断）' if interrupted else ''}：共 {len(all_articles)} 篇")
    print(f"   JSON ：{json_path}")
    if xlsx_path:
        print(f"   Excel：{xlsx_path}")
    # “已保存进度”只能在真的写出文件之后再说，以前在 except 里先嗦一句，
    # 结果往往一个文件都没落盘。
    if interrupted:
        print(f"   ↳ 中断前的 {len(all_articles)} 篇已落盘（文件名带 _partial），可稍后重跑补全；"
              "重跑会自动去重，不会重复入库。")
    if auto_ingest:
        ingest_to_kb(all_articles)
    print("\n下一步：`python3 kb.py stats` 看进度；`python3 kb.py list --unanalyzed --json`")
    print("       取批次让 agent 逐篇拆解 → `kb.py apply` 写回 → `kb.py export-html`。")
    print("=" * 64)
    return 130 if interrupted else 0


def resolve_sogou_link(url):
    """尽力把搜狗跳转链接解析为 mp.weixin.qq.com 原文（best-effort）。

    搜狗的 /link 跳转必须带上搜索页下发的 SNUID cookie，否则只会拿到一个
    空跳板页（实测 3/3 全部失败）。而且实测发现搜狗**首页只下发 SUID/ABTEST，
    只有搜索页才下发 SNUID**，所以这里用 Session 先请求一次搜索页热身。
    """
    s = requests.Session()
    s.headers.update({"User-Agent": UA, "Accept-Language": "zh-CN,zh;q=0.9"})
    warm = "https://weixin.sogou.com/weixin?type=2&query=%E5%BE%AE%E4%BF%A1"
    try:
        s.get(warm, timeout=20)
    except requests.RequestException:
        pass
    try:
        r = s.get(url, headers={"Referer": warm}, timeout=20, allow_redirects=True)
        if "mp.weixin.qq.com" in r.url:
            return r.url
        html = r.text or ""
        # 搜狗自己也有反爬：跳到 /antispider/ 时要说清楚，别让用户以为是链接问题。
        if "antispider" in r.url or "antispider" in html:
            print("⚠️  搜狗已对当前 IP 启用反爬验证（/antispider），无法解析跳转链接。")
            print("    → 请换用 mp.weixin.qq.com 原文链接，或稍后再试。")
            return url
        parts = re.findall(r"url\s*\+=\s*'([^']*)'", html)
        if parts:
            u = "".join(parts).replace("@", "").replace("&amp;", "&")
            if u.startswith("http"):
                return u
        m = re.search(r"https?://mp\.weixin\.qq\.com/s[^\"'\\<>\s]+", html)
        if m:
            return m.group(0).replace("&amp;", "&")
    except requests.RequestException:
        pass
    return url


def read_article(url, _depth=0):
    """单篇全文拓取的可复用核心，返回 (rec_or_None, status)。

    走公开 mp.weixin.qq.com/s 页，与 appmsg 后台接口是**两套限流**——这正是
    appmsg 被频控锁死时的主力备选路径：列表从别处拿（RSS/搜索/手动），正文靠它。
    run_read / run_collect_urls 共用本函数。
    """
    if not re.match(r"^https?://", url):
        return None, "invalid"
    if "weixin.sogou.com/link" in url:
        url = resolve_sogou_link(url)

    headers = {"User-Agent": UA, "Accept": "text/html,*/*", "Referer": MP_BASE + "/"}
    try:
        r = requests.get(url, headers=headers, timeout=20, allow_redirects=True)
    except requests.RequestException:
        return None, "neterror"
    html = r.text or ""
    final_url = r.url or url

    content_html = slice_content_html(html)
    content = html_to_text(content_html) if content_html else ""
    status = classify_page(html, content, final_url)

    # 分享卡片页：正文在「阅读全文」指向的原文，跟进一次（限深防环）。
    if status == "share" and _depth == 0:
        src = share_source_url(html)
        if src:
            return read_article(src, _depth=1)
    if status != "ok":
        return None, status

    title = (pick(html, r'<h1[^>]*class="rich_media_title"[^>]*>([\s\S]*?)</h1>')
             or pick(html, r'<meta property="og:title" content="([^"]*)"')
             or pick(html, r'<title>([\s\S]*?)</title>'))
    account = (pick(html, r'var nickname\s*=\s*"([^"]*)"')
               or pick(html, r'<a[^>]*id="js_name"[^>]*>([\s\S]*?)</a>')
               or pick(html, r'<meta property="og:article:author" content="([^"]*)"'))
    publish = (pick(html, r'<em[^>]*id="publish_time"[^>]*>([\s\S]*?)</em>')
               or pick(html, r'var ct\s*=\s*"(\d+)"'))
    if re.fullmatch(r"\d{10}", publish or ""):
        publish = unix_to_date(publish)
    # 用 biz+mid+idx 还原永久链接（避免搜狗 src=11 临时链重复入库）。
    perm = permanent_link(html, final_url)
    rec = {
        "id": canonical_id(perm, title),
        "account": account, "query": "", "title": title,
        "digest": pick(html, r'<meta property="og:description" content="([^"]*)"'),
        "link": perm, "cover": pick(html, r'<meta property="og:image" content="([^"]*)"'),
        "images": extract_images(content_html),
        "publishDate": publish, "content": content, "contentStatus": "ok",
        "collectedAt": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
    }
    return rec, "ok"


def run_read(url, out=None):
    """单篇全文抓取（粘链场景），输出 JSON（形态对齐 web /api/article）。"""
    if not re.match(r"^https?://", url):
        print("✗ url 格式不正确（应以 http(s):// 开头）。")
        return 1
    rec, status = read_article(url)
    if status != "ok":
        print(f"✗ 未取到正文：{STATUS_LABEL.get(status, status)}")
        if status == "blocked":
            print("  → 请隔一段时间再试，或改用浏览器打开后粘贴正文。")
        elif status == "invalid":
            print("  → 该链接已失效（搜狗的 src=11 链接寿命很短），请重新获取链接。")
        elif status == "gone":
            print("  → 文章已被删除或下架，无法恢复。")
        return 1
    if out:
        with open(out, "w", encoding="utf-8") as f:
            json.dump(rec, f, ensure_ascii=False, indent=2)
        print(f"✓ 已写入 {out}（{len(rec['content'])} 字 / {len(rec['images'])} 图）。")
    else:
        print(json.dumps(rec, ensure_ascii=False, indent=2))
    return 0


# ═══════════════════════════════════════════════════════════
#  子命令：collect-urls（appmsg 被频控时的备选路径：URL/文件/RSS → 逐篇 read）
# ═══════════════════════════════════════════════════════════

def _extract_mp_links(text):
    """从任意文本/RSS/Atom 中抽出 mp.weixin.qq.com 文章链接（保序去重）。"""
    text = (text or "").replace("&amp;", "&")
    urls, seen = [], set()
    for m in re.finditer(r"https?://mp\.weixin\.qq\.com/s[^\s\"'<>\\)\]]+", text):
        u = m.group(0).rstrip(".,;")
        if u not in seen:
            seen.add(u)
            urls.append(u)
    return urls


def load_urls(source):
    """把一个来源解析成 URL 列表。source 可以是：
      - 本地文件：.txt（每行一个，# 开头为注释）/ .json（数组或 {articles:[{link}]}）/ .opml/.xml
      - RSS/Atom 订阅链接（http…）→ 拉取并抽出其中的 mp 链接
      - 直接就是一个 mp 文章链接
    """
    # RSS/Atom 或单个网址
    if re.match(r"^https?://", source):
        if "mp.weixin.qq.com/s" in source:
            return [source]
        try:
            r = requests.get(source, headers={"User-Agent": UA}, timeout=25)
            links = _extract_mp_links(r.text)
            print(f"   ↓ 从订阅源解析出 {len(links)} 个公众号文章链接。")
            return links
        except requests.RequestException as e:
            print(f"✗ 拉取订阅源失败：{e}")
            return []
    # 本地文件
    if os.path.exists(source):
        raw = open(source, "r", encoding="utf-8").read()
        if source.endswith(".json"):
            try:
                data = json.loads(raw)
                items = data.get("articles") if isinstance(data, dict) else data
                out = []
                for it in items or []:
                    u = it.get("link") or it.get("url") if isinstance(it, dict) else it
                    if u:
                        out.append(u)
                return out
            except (ValueError, AttributeError):
                pass
        # .txt / .opml / .xml / 其它：先按行取，再兵底正则抽 mp 链接
        lines = [ln.strip() for ln in raw.splitlines()
                 if ln.strip() and not ln.strip().startswith("#")]
        if lines and all(re.match(r"^https?://", ln) for ln in lines):
            return lines
        return _extract_mp_links(raw)
    print(f"✗ 来源无法识别（既不是 http(s) 链接也不是存在的文件）：{source}")
    return []


def run_collect_urls(sources, auto_ingest):
    """备选采集：从 URL / 文件 / RSS 拿到一批链接，逐篇走 read 拓全文。

    这条路**不碰 appmsg**，所以 appmsg 被频控锁死的当天依然能用。
    """
    print("=" * 64)
    print("  墨摘 · 备选采集（URL/文件/RSS → 逐篇全文，不经 appmsg）")
    print("=" * 64)
    print("ℹ️  这是**备选**路径：首选仍是 collect（appmsg，最准确），仅在它被频控拿不到时才用这条。")
    print("    本路径的文章完整性只取决于你给的来源（可能漏文章/串同名号），请知悉。")

    urls = []
    for src in sources:
        for u in load_urls(src):
            if u not in urls:
                urls.append(u)
    if not urls:
        print("✗ 没有拿到任何文章链接。")
        return 1
    print(f"\n📎 共 {len(urls)} 个链接，开始逐篇拓全文…")

    collected, reasons = [], {}
    for i, u in enumerate(urls, 1):
        rec, status = read_article(u)
        if status == "ok" and rec:
            collected.append(rec)
            print(f"   [{len(collected)}] {rec['account']} · {rec['publishDate']} · "
                  f"{len(rec['content'])}字/{len(rec['images'])}图  {rec['title'][:30]}")
        else:
            reasons[status] = reasons.get(status, 0) + 1
            print(f"   [✗] {STATUS_LABEL.get(status, status)}  <- …{u[-22:]}")
        polite_sleep(CONTENT_DELAY_MIN, CONTENT_DELAY_MAX)

    n = len(urls)
    print(f"\n   ↳ 成功 {len(collected)}/{n}（空率 {(n - len(collected)) / n * 100:.0f}%）")
    for st, c in sorted(reasons.items(), key=lambda kv: -kv[1]):
        print(f"      · {STATUS_LABEL.get(st, st)} × {c}")
    if not collected:
        return 1

    accounts = sorted({r["account"] for r in collected if r["account"]})
    json_path = save_json(collected, accounts, "")
    xlsx_path = save_excel(collected)
    print("\n" + "=" * 64)
    print(f"✅ 完成：共 {len(collected)} 篇")
    print(f"   JSON ：{json_path}")
    if xlsx_path:
        print(f"   Excel：{xlsx_path}")
    if auto_ingest:
        ingest_to_kb(collected)
    print("=" * 64)
    return 0


def run_whoami():
    """用一次轻量 searchbiz 校验 token/cookie 是否有效。"""
    token, cookie = load_credentials()
    if not token or not cookie:
        print("✗ 未检测到 token/cookie。")
        print_credential_guide()
        return 1
    print("· 正在用一次轻量搜索校验登录态…")
    session = make_session(cookie)
    try:
        api_get(session, token, "/cgi-bin/searchbiz",
                {"action": "search_biz", "begin": "0", "count": "1", "query": "微信"})
        print(f"✓ 登录态有效。token=…{token[-4:]}，cookie 长度={len(cookie)}。可以开始采集。")
        return 0
    except AuthError as e:
        print(f"✗ 凭证失效：{e}")
        print_credential_guide()
        return 1
    except FreqControlError as e:
        print(f"⏸ 命中频控（但凭证看起来有效）：{e}")
        return 0
    except ApiError as e:
        print(f"⚠️ 接口异常：{e}")
        return 1


# ════════════════════════════════════════════════════════════════
#  入口
# ════════════════════════════════════════════════════════════════

def build_parser():
    p = argparse.ArgumentParser(
        prog="wechat_collector.py",
        description="墨摘 · 微信公众号采集（cookie+token 后台接口，P0 全文挖掘）")
    sub = p.add_subparsers(dest="cmd")

    sp = sub.add_parser("collect", help="按公众号名称采集文章（默认抓正文 + 自动入库）")
    sp.add_argument("accounts", nargs="+", help="公众号名称（可多个，精确名命中率更高）")
    sp.add_argument("--since", default=DATE_FILTER_THRESHOLD, help="仅采此日期(含)之后；传空串则不限")
    sp.add_argument("--count", type=int, default=TARGET_COUNT, help="每号目标篇数")
    sp.add_argument("--no-content", action="store_true", help="只采列表元数据，不抓正文")
    sp.add_argument("--no-kb", action="store_true", help="不自动入知识库")
    sp.add_argument("--resume", action="store_true",
                    help="断点续采：从上次存的 begin 游标接着翻，不重复请求已拉页（拉大号历史、省频控配额）")
    _add_freq_args(sp)

    sp = sub.add_parser("read", help="抓取单篇文章全文（粘链场景）")
    sp.add_argument("url")
    sp.add_argument("--out", help="写入 JSON 文件路径（默认打印到 stdout）")

    sp = sub.add_parser("collect-urls",
                        help="备选采集：从 URL/文件/RSS 批量拓全文（appmsg 被频控时用）")
    sp.add_argument("sources", nargs="+",
                    help="mp 文章链接 / 本地文件(.txt/.json/.opml) / RSS或Atom订阅链接（可混合多个）")
    sp.add_argument("--no-kb", action="store_true", help="不自动入知识库")

    sub.add_parser("whoami", help="校验 token/cookie 是否有效")

    sp = sub.add_parser("search", help="按关键词搜索全网公众号文章（⚠ 接口未经验证）")
    sp.add_argument("keyword", help="搜索关键词")
    sp.add_argument("--count", type=int, default=10, help="目标篇数")
    sp.add_argument("--no-content", action="store_true", help="只采列表元数据，不抓正文")
    sp.add_argument("--no-kb", action="store_true", help="不自动入知识库")
    _add_freq_args(sp)

    return p


def _add_freq_args(sp):
    """频控节奏参数：不同账号/IP 的频控尺度差很多，开放给用户自调。"""
    sp.add_argument("--freq-cooldown", type=int, default=FREQ_COOLDOWN,
                    help=f"命中频控(ret=200013)后的首次冷却秒数（默认 {FREQ_COOLDOWN}，逐次翻倍）")
    sp.add_argument("--freq-retries", type=int, default=MAX_FREQ_RETRIES,
                    help=f"频控重试次数（默认 {MAX_FREQ_RETRIES}；设 0 则一撞频控就落盘退出）")


def _apply_freq_args(args):
    """把 CLI 传的频控参数写回模块级常量。"""
    global FREQ_COOLDOWN, MAX_FREQ_RETRIES
    if getattr(args, "freq_cooldown", None) is not None:
        FREQ_COOLDOWN = max(0, args.freq_cooldown)
    if getattr(args, "freq_retries", None) is not None:
        MAX_FREQ_RETRIES = max(0, args.freq_retries)


def main():
    parser = build_parser()
    args = parser.parse_args()
    _apply_freq_args(args)

    if args.cmd == "collect":
        return run_collect(args.accounts, args.since.strip(), args.count,
                           not args.no_content, not args.no_kb, args.resume)
    if args.cmd == "read":
        return run_read(args.url, args.out)
    if args.cmd == "collect-urls":
        return run_collect_urls(args.sources, not args.no_kb)
    if args.cmd == "whoami":
        return run_whoami()
    if args.cmd == "search":
        return run_search(args.keyword, args.count, not args.no_content, not args.no_kb)

    # 无子命令 → 旧用法：读文件顶部默认配置
    return run_collect(ACCOUNTS, DATE_FILTER_THRESHOLD, TARGET_COUNT, FETCH_CONTENT, AUTO_INGEST, False)


if __name__ == "__main__":
    try:
        sys.exit(main() or 0)
    except KeyboardInterrupt:
        print("\n已手动中断。")
        sys.exit(130)
